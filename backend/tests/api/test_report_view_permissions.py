"""7e R1 报表视图权限的路由层契约（报表视图设计第 4 节、5.1）。

真实路由 + 迁移 seed 的真实角色（``flow_users``）：

- 工作进度：``require_any_permission(read, read_own)``；只持 ``read_own`` 的 PR 只看到本人那一行
  （service 层 ``resolve_pr_scope`` 强制）；区间参数 ``preset`` / ``date_from`` / ``date_to``，
  旧参数 ``month`` 换算成整月；F3 聚合金额运营与管理员相同
- 发文进度 4 个接口：``require_any_permission(read, read_own)``；只持 ``read_own`` 的 PR 只看到
  本人的单（``PrScope`` 下推到 SQL，摘要 / 卡片 total 也只数本人）
- BI 三个端点挂 ``report.bi:read``（不再挂 ``report.production:read``）；推广汇总 4 个金额过 F3
- 导出：路由仍要 ``report.export:read``，service 再按报表类型判对应读权限（缺 → 403）
- ``GET /api/reports/capabilities``：只要登录，6 个真实角色逐字段期望
"""

from __future__ import annotations

import io
from collections.abc import AsyncIterator
from datetime import timedelta
from decimal import Decimal
from typing import Any
from uuid import uuid4

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.tenancy import tenant_id_ctx
from app.modules.auth.models import Permission, UserPermissionOverride
from app.modules.promotion.urge_calculator import get_today

pytestmark = [pytest.mark.api, pytest.mark.asyncio]

# 场景放在上个月：整月都在过去，日期与「今天」无关的计数稳定
_TODAY = get_today()
_M_HI = _TODAY.replace(day=1) - timedelta(days=1)
_M_LO = _M_HI.replace(day=1)
_MONTH = f"{_M_LO:%Y-%m}"
_WP = "/api/reports/work-progress"


async def _call(session: AsyncSession, user: Any, method: str, path: str, json: Any = None) -> Any:
    """真实路由（含 schema 校验）+ 用户的有效权限。"""
    from app.core.db import get_session
    from app.main import app
    from app.modules.auth.deps import get_current_perms, get_current_user_active
    from app.modules.auth.service import AuthService

    async def _session_override() -> AsyncIterator[AsyncSession]:
        yield session

    perms = await AuthService(session).load_effective_permissions(user.id)
    try:
        app.dependency_overrides[get_session] = _session_override
        app.dependency_overrides[get_current_user_active] = lambda: user
        app.dependency_overrides[get_current_perms] = lambda: perms
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
            return await c.request(method, path, json=json)
    finally:
        for dep in (get_session, get_current_user_active, get_current_perms):
            app.dependency_overrides.pop(dep, None)


@pytest.fixture
def tenant_ctx(tenant_a: Any) -> Any:
    token = tenant_id_ctx.set(tenant_a.id)
    yield
    tenant_id_ctx.reset(token)


@pytest.fixture
async def seeded(
    flow_users: Any, product_factory: Any, blogger_factory: Any, promotion_factory: Any
) -> Any:
    """上个月：pr 两单、pr2 一单、未分配一单，都带样品成本（工作进度的 ``cost``）。"""
    style = await product_factory.style(style_code=f"RV{uuid4().hex[:8]}")
    blogger = await blogger_factory.blogger()

    async def promo(**kw: Any) -> None:
        await promotion_factory.promotion(
            style=style,
            blogger=blogger,
            internal_code=f"RV{uuid4().hex[:10].upper()}",
            cooperation_date=_M_LO + timedelta(days=2),
            **kw,
        )

    await promo(pr=flow_users.pr, cost_snapshot=Decimal("30.00"), publish_status="已发布")
    await promo(pr=flow_users.pr, cost_snapshot=Decimal("12.50"))
    await promo(pr=flow_users.pr2, cost_snapshot=Decimal("40.00"))
    await promo(cost_snapshot=Decimal("7.00"))
    return flow_users


@pytest.fixture
async def goods_seeded(
    session: AsyncSession,
    seeded: Any,
    product_factory: Any,
    blogger_factory: Any,
    promotion_factory: Any,
) -> Any:
    """在 ``seeded`` 之上：pr 一单归属商品甲、pr2 一单归属商品乙（单品卡片 / 周明细的只看本人）。"""
    from app.modules.product.goods_models import GoodsMain

    style = await product_factory.style(style_code=f"RG{uuid4().hex[:8]}")
    blogger = await blogger_factory.blogger()
    made: list[Any] = []
    for user in (seeded.pr, seeded.pr2):
        goods = GoodsMain(
            tenant_id=seeded.tenant.id,
            goods_code=f"RG{uuid4().hex[:8]}",
            goods_title=f"商品{uuid4().hex[:4]}",
        )
        session.add(goods)
        await session.flush()
        await promotion_factory.promotion(
            style=style,
            blogger=blogger,
            pr=user,
            goods_main_id=goods.id,
            internal_code=f"RG{uuid4().hex[:10].upper()}",
            cooperation_date=_M_LO + timedelta(days=9),
        )
        made.append(goods)
    await session.flush()
    return seeded, made[0], made[1]


@pytest.mark.usefixtures("tenant_ctx")
class TestWorkProgress:
    async def test_pr_sees_only_self(self, session: AsyncSession, seeded: Any) -> None:
        resp = await _call(session, seeded.pr, "GET", f"{_WP}?month={_MONTH}")
        assert resp.status_code == 200, resp.text
        rows = resp.json()
        assert [r["pr_id"] for r in rows] == [str(seeded.pr.id)]
        assert rows[0]["quote_count"] == 2

    async def test_finance_sees_all(self, session: AsyncSession, seeded: Any) -> None:
        resp = await _call(session, seeded.finance, "GET", f"{_WP}?month={_MONTH}")
        assert resp.status_code == 200, resp.text
        assert {r["pr_id"] for r in resp.json()} == {str(seeded.pr.id), str(seeded.pr2.id), None}

    async def test_month_equals_custom_range(self, session: AsyncSession, seeded: Any) -> None:
        by_month = await _call(session, seeded.pr_manager, "GET", f"{_WP}?month={_MONTH}")
        by_range = await _call(
            session,
            seeded.pr_manager,
            "GET",
            f"{_WP}?preset=custom&date_from={_M_LO}&date_to={_M_HI}",
        )
        assert by_month.status_code == by_range.status_code == 200
        assert by_month.json() == by_range.json()
        # 场景有效性：三行（pr / pr2 / 未分配）
        assert len(by_month.json()) == 3

    async def test_default_preset_is_this_month(self, session: AsyncSession, seeded: Any) -> None:
        bare = await _call(session, seeded.pr_manager, "GET", _WP)
        this_month = await _call(session, seeded.pr_manager, "GET", f"{_WP}?preset=this_month")
        assert bare.status_code == this_month.status_code == 200
        assert bare.json() == this_month.json()

    async def test_operations_aggregate_cost_equals_admin(
        self, session: AsyncSession, seeded: Any
    ) -> None:
        """F3：运营读不到逐单 ``cost_snapshot``，聚合的 ``cost`` / ``cpl`` 照样与管理员相同。"""
        ops = await _call(session, seeded.operations, "GET", f"{_WP}?month={_MONTH}")
        admin = await _call(session, seeded.admin, "GET", f"{_WP}?month={_MONTH}")
        assert ops.status_code == admin.status_code == 200
        assert ops.json() == admin.json()
        costs = {r["pr_id"]: r["cost"] for r in ops.json()}
        assert costs[str(seeded.pr.id)] == "42.50"

    async def test_warehouse_forbidden(self, session: AsyncSession, seeded: Any) -> None:
        resp = await _call(session, seeded.warehouse, "GET", f"{_WP}?month={_MONTH}")
        assert (resp.status_code, resp.json()["code"]) == (403, "PERMISSION_DENIED")

    @pytest.mark.parametrize(
        ("query", "code"),
        [
            ("preset=foo", "REPORT_INVALID_TIME_PRESET"),
            (
                f"preset=custom&date_from={_M_LO}&date_to={_M_LO + timedelta(days=367)}",
                "REPORT_INVALID_TIME_RANGE",
            ),
            (f"preset=custom&date_from={_M_HI}&date_to={_M_LO}", "REPORT_INVALID_TIME_RANGE"),
            ("month=2026-13", "REPORT_INVALID_TIME_RANGE"),
        ],
        ids=["bad_preset", "span_367", "from_after_to", "bad_month"],
    )
    async def test_invalid_range_422(
        self, session: AsyncSession, flow_users: Any, query: str, code: str
    ) -> None:
        resp = await _call(session, flow_users.pr_manager, "GET", f"{_WP}?{query}")
        assert (resp.status_code, resp.json()["code"]) == (422, code)

    async def test_span_366_ok(self, session: AsyncSession, flow_users: Any) -> None:
        hi = _M_LO + timedelta(days=366)
        resp = await _call(
            session,
            flow_users.pr_manager,
            "GET",
            f"{_WP}?preset=custom&date_from={_M_LO}&date_to={hi}",
        )
        assert resp.status_code == 200, resp.text

    @pytest.mark.parametrize(
        "season_q",
        ["".join(f"&season=s{i}" for i in range(21)), f"&season={'x' * 65}"],
        ids=["21_seasons", "65_chars"],
    )
    async def test_invalid_season_422(
        self, session: AsyncSession, flow_users: Any, season_q: str
    ) -> None:
        resp = await _call(session, flow_users.pr_manager, "GET", f"{_WP}?month={_MONTH}{season_q}")
        assert (resp.status_code, resp.json()["code"]) == (422, "VALIDATION_ERROR")

    async def test_season_keeps_only_matching_goods(
        self,
        session: AsyncSession,
        seeded: Any,
        promotion_factory: Any,
        product_factory: Any,
        blogger_factory: Any,
    ) -> None:
        """带季节只统计归属商品季节命中的单：seeded 的单都没归属商品，选了季节就不算。"""
        from app.modules.product.goods_models import GoodsMain

        tenant = seeded.tenant
        goods = GoodsMain(
            tenant_id=tenant.id,
            goods_code=f"WS{uuid4().hex[:8]}",
            goods_title="季节商品",
            season="2026秋",
        )
        session.add(goods)
        await session.flush()
        style = await product_factory.style(style_code=f"WS{uuid4().hex[:8]}")
        blogger = await blogger_factory.blogger()
        await promotion_factory.promotion(
            style=style,
            blogger=blogger,
            pr=seeded.pr,
            goods_main_id=goods.id,
            internal_code=f"WS{uuid4().hex[:10].upper()}",
            cooperation_date=_M_LO + timedelta(days=3),
            quote_amount=Decimal("88.00"),
        )
        await session.flush()

        bare = await _call(session, seeded.finance, "GET", f"{_WP}?month={_MONTH}")
        autumn = await _call(session, seeded.finance, "GET", f"{_WP}?month={_MONTH}&season=2026秋")
        other = await _call(session, seeded.finance, "GET", f"{_WP}?month={_MONTH}&season=春夏")
        assert bare.status_code == autumn.status_code == other.status_code == 200
        # 场景有效性：不带季节时 pr 有 3 单（seeded 2 + 本单）、共 3 行
        assert {r["pr_id"]: r["quote_count"] for r in bare.json()}[str(seeded.pr.id)] == 3
        assert [(r["pr_id"], r["quote_count"], r["quote_amount"]) for r in autumn.json()] == [
            (str(seeded.pr.id), 1, "88.00")
        ]
        assert other.json() == []

    async def test_goods_route_defaults(self, session: AsyncSession, seeded: Any) -> None:
        """单品卡片（7e R5）：默认 ``limit`` 12、``offset`` 0；seeded 的单都没归属商品 → 只有信封里的未归属。"""
        resp = await _call(
            session,
            seeded.finance,
            "GET",
            f"{_WP}/goods?preset=custom&date_from={_M_LO}&date_to={_M_HI}",
        )
        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert (body["items"], body["total"], body["limit"], body["offset"]) == ([], 0, 12, 0)
        assert (body["date_from"], body["date_to"]) == (str(_M_LO), str(_M_HI))
        assert body["unassigned"]["display_name"] == "未归属商品"
        assert body["unassigned"]["quote_count"] == 4
        assert "goods_code" not in body["unassigned"]

    async def test_weeks_route(self, session: AsyncSession, seeded: Any) -> None:
        """周明细（7e R5）：区间原样回传；各周约稿量之和 = PR 表合计；行起止都在区间内。"""
        query = f"preset=custom&date_from={_M_LO}&date_to={_M_HI}"
        resp = await _call(session, seeded.finance, "GET", f"{_WP}/weeks?{query}")
        table = await _call(session, seeded.finance, "GET", f"{_WP}?{query}")
        assert resp.status_code == table.status_code == 200, resp.text
        body = resp.json()
        assert (body["date_from"], body["date_to"]) == (str(_M_LO), str(_M_HI))
        assert body["items"]
        assert sum(r["quote_count"] for r in body["items"]) == sum(
            r["quote_count"] for r in table.json()
        )
        for r in body["items"]:
            assert str(_M_LO) <= r["date_from"] <= r["date_to"] <= str(_M_HI)
            assert r["week_start"] <= r["date_from"]

    async def test_goods_pr_sees_only_self(self, session: AsyncSession, goods_seeded: Any) -> None:
        """只持 ``read_own`` 的 PR：卡片只有本人的商品与本人的未归属单；财务看全部。"""
        users, own_goods, other_goods = goods_seeded
        query = f"{_WP}/goods?preset=custom&date_from={_M_LO}&date_to={_M_HI}"
        own = await _call(session, users.pr, "GET", query)
        everyone = await _call(session, users.finance, "GET", query)
        assert own.status_code == everyone.status_code == 200, own.text
        own_body, all_body = own.json(), everyone.json()
        assert [(r["goods_id"], r["quote_count"]) for r in own_body["items"]] == [
            (str(own_goods.id), 1)
        ]
        assert str(other_goods.id) not in {r["goods_id"] for r in own_body["items"]}
        assert own_body["total"] == 1
        assert own_body["unassigned"]["quote_count"] == 2  # seeded 里 pr 的两单
        # 场景有效性：财务能看到别人的商品与全部未归属单
        assert {r["goods_id"] for r in all_body["items"]} == {
            str(own_goods.id),
            str(other_goods.id),
        }
        assert all_body["unassigned"]["quote_count"] == 4

    async def test_weeks_pr_sees_only_self(self, session: AsyncSession, goods_seeded: Any) -> None:
        """只持 ``read_own`` 的 PR：各周之和 = 本人那一行 PR 表；财务各周之和 = 全部 PR 表合计。"""
        users, _own_goods, _other_goods = goods_seeded
        query = f"preset=custom&date_from={_M_LO}&date_to={_M_HI}"

        async def totals(user: Any) -> tuple[int, int]:
            weeks = await _call(session, user, "GET", f"{_WP}/weeks?{query}")
            table = await _call(session, user, "GET", f"{_WP}?{query}")
            assert weeks.status_code == table.status_code == 200, weeks.text
            return (
                sum(r["quote_count"] for r in weeks.json()["items"]),
                sum(r["quote_count"] for r in table.json()),
            )

        assert await totals(users.pr) == (3, 3)  # seeded 2 + 本人商品 1
        assert await totals(users.finance) == (6, 6)

    async def test_weeks_warehouse_forbidden(self, session: AsyncSession, flow_users: Any) -> None:
        resp = await _call(session, flow_users.warehouse, "GET", f"{_WP}/weeks")
        assert (resp.status_code, resp.json()["code"]) == (403, "PERMISSION_DENIED")

    @pytest.mark.parametrize("query", ["limit=101", "limit=0", "offset=-1"])
    async def test_goods_paging_422(
        self, session: AsyncSession, flow_users: Any, query: str
    ) -> None:
        resp = await _call(session, flow_users.pr_manager, "GET", f"{_WP}/goods?{query}")
        assert resp.status_code == 422, resp.text

    async def test_goods_limit_100_ok(self, session: AsyncSession, flow_users: Any) -> None:
        resp = await _call(session, flow_users.pr_manager, "GET", f"{_WP}/goods?limit=100")
        assert resp.status_code == 200, resp.text

    async def test_goods_warehouse_forbidden(self, session: AsyncSession, flow_users: Any) -> None:
        resp = await _call(session, flow_users.warehouse, "GET", f"{_WP}/goods")
        assert (resp.status_code, resp.json()["code"]) == (403, "PERMISSION_DENIED")

    async def test_goods_requires_login(self, session: AsyncSession) -> None:
        from app.core.db import get_session
        from app.main import app

        async def _session_override() -> AsyncIterator[AsyncSession]:
            yield session

        app.dependency_overrides[get_session] = _session_override
        try:
            async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
                resp = await c.get(f"{_WP}/goods")
        finally:
            app.dependency_overrides.pop(get_session, None)
        assert resp.status_code == 401

    @pytest.mark.parametrize("role", ["pr", "finance", "pr_manager"])
    async def test_season_options_open_to_work_progress_readers(
        self, session: AsyncSession, flow_users: Any, role: str
    ) -> None:
        """季节下拉只回季节名：持工作进度 read / read_own 就能拿（7e R5，原来只认投产读权限）。"""
        resp = await _call(session, getattr(flow_users, role), "GET", "/api/reports/season-options")
        assert resp.status_code == 200, resp.text
        assert isinstance(resp.json()["items"], list)

    async def test_season_options_warehouse_forbidden(
        self, session: AsyncSession, flow_users: Any
    ) -> None:
        resp = await _call(session, flow_users.warehouse, "GET", "/api/reports/season-options")
        assert (resp.status_code, resp.json()["code"]) == (403, "PERMISSION_DENIED")


_PP = "/api/reports/publish-progress"
_PP_RANGE = f"preset=custom&date_from={_M_LO}&date_to={_M_HI}"


@pytest.fixture
async def pp_seeded(
    flow_users: Any, product_factory: Any, blogger_factory: Any, promotion_factory: Any
) -> Any:
    """上个月：款 A 上 pr 两单（一单已发布）、pr2 一单、未分配一单；款 B 只有 pr2 一单。

    返回 ``(users, style_a, style_b, 非 pr 的单 id)``。
    """
    style_a = await product_factory.style(style_code=f"PA{uuid4().hex[:8]}")
    style_b = await product_factory.style(style_code=f"PB{uuid4().hex[:8]}")
    blogger = await blogger_factory.blogger()

    async def promo(style: Any, day: int, **kw: Any) -> Any:
        return await promotion_factory.promotion(
            style=style,
            blogger=blogger,
            internal_code=f"PP{uuid4().hex[:10].upper()}",
            cooperation_date=_M_LO + timedelta(days=day),
            **kw,
        )

    pr = flow_users.pr
    await promo(
        style_a,
        2,
        pr=pr,
        quote_amount=Decimal("300.00"),
        publish_status="已发布",
        like_count=120,
    )
    await promo(style_a, 20, pr=pr, quote_amount=Decimal("150.00"))
    others = [
        await promo(
            style_a,
            3,
            pr=flow_users.pr2,
            quote_amount=Decimal("500.00"),
            publish_status="已发布",
            like_count=900,
        ),
        await promo(style_a, 21, quote_amount=Decimal("80.00")),
        await promo(style_b, 4, pr=flow_users.pr2, quote_amount=Decimal("60.00")),
    ]
    return flow_users, style_a, style_b, [p.id for p in others]


def _pp_paths(style_a: Any, style_b: Any) -> dict[str, str]:
    return {
        "summary": f"{_PP}/summary?{_PP_RANGE}",
        "cards": f"{_PP}/cards?{_PP_RANGE}",
        "by_pr_a": f"{_PP}/styles/{style_a.id}/by-pr?{_PP_RANGE}",
        "by_pr_b": f"{_PP}/styles/{style_b.id}/by-pr?{_PP_RANGE}",
        "by_time_a": f"{_PP}/styles/{style_a.id}/by-time?{_PP_RANGE}",
    }


async def _pp_all(session: AsyncSession, user: Any, paths: dict[str, str]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for name, path in paths.items():
        resp = await _call(session, user, "GET", path)
        assert resp.status_code == 200, (name, resp.text)
        out[name] = resp.json()
    return out


@pytest.mark.usefixtures("tenant_ctx")
class TestPublishProgress:
    async def test_pr_sees_only_self(self, session: AsyncSession, pp_seeded: Any) -> None:
        """pr 的 4 个接口 = 只有 pr 的单时主管看到的数字（逐字相等）。"""
        users, style_a, style_b, others = pp_seeded
        paths = _pp_paths(style_a, style_b)
        mine = await _pp_all(session, users.pr, paths)
        full = await _pp_all(session, users.pr_manager, paths)

        # 场景有效性：主管看到的比 pr 多（别人的单确实在区间里）
        assert full["summary"]["quote_count"] == 5
        assert full["cards"]["total"] == 2
        assert len(full["by_pr_a"]) == 3
        assert sum(p["quote_count"] for p in full["by_time_a"]) == 4

        # pr 只剩本人：数字非零，且只有本人那一行
        assert mine["summary"]["quote_count"] == 2
        assert mine["summary"]["publish_count"] == 1
        assert mine["summary"]["quote_amount"] == "450.00"
        assert mine["summary"]["like_count"] > 0
        assert mine["cards"]["total"] == 1
        assert [c["style_id"] for c in mine["cards"]["items"]] == [str(style_a.id)]
        assert [r["pr_id"] for r in mine["by_pr_a"]] == [str(users.pr.id)]
        assert mine["by_pr_b"] == []
        assert sum(p["quote_count"] for p in mine["by_time_a"]) == 2

        # 停用别人的单后，主管看到的 = pr 之前看到的
        await session.execute(
            text("UPDATE promotion SET is_active = false WHERE id = ANY(CAST(:ids AS uuid[]))"),
            {"ids": others},
        )
        only_mine = await _pp_all(session, users.pr_manager, paths)
        assert only_mine == mine

    async def test_admin_sees_all(self, session: AsyncSession, pp_seeded: Any) -> None:
        users, style_a, _, _ = pp_seeded
        resp = await _call(session, users.admin, "GET", f"{_PP}/summary?{_PP_RANGE}")
        assert resp.status_code == 200, resp.text
        assert resp.json()["quote_count"] == 5

    @pytest.mark.parametrize("role", ["finance", "warehouse"])
    async def test_forbidden(self, session: AsyncSession, pp_seeded: Any, role: str) -> None:
        users, style_a, style_b, _ = pp_seeded
        for path in _pp_paths(style_a, style_b).values():
            resp = await _call(session, getattr(users, role), "GET", path)
            assert (resp.status_code, resp.json()["code"]) == (403, "PERMISSION_DENIED"), path


async def _override(session: AsyncSession, user: Any, scope: str, effect: str) -> None:
    """给用户造一条个人授予 / 撤销（scope 须已入册）。"""
    perm = (await session.execute(select(Permission).where(Permission.scope == scope))).scalar_one()
    session.add(
        UserPermissionOverride(
            tenant_id=user.tenant_id, user_id=user.id, permission_id=perm.id, effect=effect
        )
    )
    await session.flush()


_BI = "/api/reports/bi"
_BI_RANGE = f"preset=custom&date_from={_M_LO}&date_to={_M_HI}"
_PROD = f"/api/reports/production?{_BI_RANGE}"


@pytest.fixture
async def bi_seeded(
    flow_users: Any, product_factory: Any, blogger_factory: Any, promotion_factory: Any
) -> Any:
    """上个月：已发布 / 未发布 / 已取消各一单，推广汇总 4 个金额都非零。"""
    style = await product_factory.style(style_code=f"BI{uuid4().hex[:8]}")
    blogger = await blogger_factory.blogger()
    for i, (status, amount) in enumerate(
        [("已发布", "300.00"), ("未发布", "150.00"), ("已取消", "80.00")]
    ):
        await promotion_factory.promotion(
            style=style,
            blogger=blogger,
            pr=flow_users.pr,
            internal_code=f"BI{uuid4().hex[:10].upper()}",
            cooperation_date=_M_LO + timedelta(days=i + 1),
            quote_amount=Decimal(amount),
            publish_status=status,
        )
    return flow_users


@pytest.mark.usefixtures("tenant_ctx")
class TestBi:
    async def test_finance_bi_and_layout_but_not_production(
        self, session: AsyncSession, flow_users: Any
    ) -> None:
        fin = flow_users.finance
        bi = await _call(session, fin, "GET", f"{_BI}?{_BI_RANGE}")
        assert bi.status_code == 200, bi.text
        get_layout = await _call(session, fin, "GET", f"{_BI}/layout")
        assert get_layout.status_code == 200, get_layout.text
        put_layout = await _call(session, fin, "PUT", f"{_BI}/layout", json={"cards": ["x"]})
        assert put_layout.status_code == 200, put_layout.text
        assert (await _call(session, fin, "GET", f"{_BI}/layout")).json() == {"cards": ["x"]}
        # BI 拆出来后，给财务 BI 不再连带放开投产页
        prod = await _call(session, fin, "GET", _PROD)
        assert (prod.status_code, prod.json()["code"]) == (403, "PERMISSION_DENIED")

    async def test_production_read_alone_is_not_bi(
        self, session: AsyncSession, flow_users: Any
    ) -> None:
        wh = flow_users.warehouse
        await _override(session, wh, "report.production:read", "grant")
        # 场景有效性：授予生效（投产页能进）
        assert (await _call(session, wh, "GET", _PROD)).status_code == 200
        for method, path, body in [
            ("GET", f"{_BI}?{_BI_RANGE}", None),
            ("GET", f"{_BI}/layout", None),
            ("PUT", f"{_BI}/layout", {"cards": ["x"]}),
        ]:
            resp = await _call(session, wh, method, path, json=body)
            assert (resp.status_code, resp.json()["code"]) == (403, "PERMISSION_DENIED"), path

    @pytest.mark.parametrize("role", ["pr", "warehouse"])
    async def test_forbidden(self, session: AsyncSession, flow_users: Any, role: str) -> None:
        resp = await _call(session, getattr(flow_users, role), "GET", f"{_BI}?{_BI_RANGE}")
        assert (resp.status_code, resp.json()["code"]) == (403, "PERMISSION_DENIED")

    async def test_operations_promo_amounts_equal_admin(
        self, session: AsyncSession, bi_seeded: Any
    ) -> None:
        """F3：运营读不到逐单 ``quote_amount``，BI 推广汇总金额照样与管理员相同。"""
        ops = await _call(session, bi_seeded.operations, "GET", f"{_BI}?{_BI_RANGE}")
        admin = await _call(session, bi_seeded.admin, "GET", f"{_BI}?{_BI_RANGE}")
        assert ops.status_code == admin.status_code == 200
        assert ops.json() == admin.json()
        summary = ops.json()["promotion_summary"]
        # 场景有效性：4 个金额都非零
        assert (
            summary["commission_amount"],
            summary["published_spend"],
            summary["unpublished_spend"],
            summary["cancelled_amount"],
        ) == ("530.00", "300.00", "150.00", "80.00")


def _export(kind: str) -> str:
    return f"/api/reports/{kind}/export?{_BI_RANGE}"


@pytest.mark.usefixtures("tenant_ctx")
class TestExport:
    @pytest.mark.parametrize(
        ("kind", "status"),
        [("work-progress", 200), ("production", 403), ("store-daily", 403)],
    )
    async def test_finance_by_report_type(
        self, session: AsyncSession, flow_users: Any, kind: str, status: int
    ) -> None:
        """财务持 ``report.export:read`` + 工作进度 read：能导工作进度，投产 / 店铺 403。"""
        resp = await _call(session, flow_users.finance, "GET", _export(kind))
        assert resp.status_code == status, resp.text
        if status == 403:
            assert resp.json()["code"] == "PERMISSION_DENIED"
            assert resp.json()["details"]["required_action"] == "read"
        else:
            assert resp.headers["content-type"].startswith(
                "application/vnd.openxmlformats-officedocument.spreadsheetml"
            )

    async def test_work_progress_three_sheets_and_season(
        self,
        session: AsyncSession,
        seeded: Any,
        product_factory: Any,
        blogger_factory: Any,
        promotion_factory: Any,
    ) -> None:
        """工作进度导出（7e R5）：三个 sheet；带季节只剩命中的单；单品末行未归属、不导商品编码。"""
        from app.modules.product.goods_models import GoodsMain

        goods = GoodsMain(
            tenant_id=seeded.tenant.id,
            goods_code=f"XS{uuid4().hex[:8]}",
            goods_title="季节导出商品",
            season="2026秋",
        )
        session.add(goods)
        await session.flush()
        await promotion_factory.promotion(
            style=await product_factory.style(style_code=f"XS{uuid4().hex[:8]}"),
            blogger=await blogger_factory.blogger(),
            pr=seeded.pr,
            goods_main_id=goods.id,
            internal_code=f"XS{uuid4().hex[:10].upper()}",
            cooperation_date=_M_LO + timedelta(days=3),
            quote_amount=Decimal("66.00"),
        )
        await session.flush()

        async def sheets(query: str) -> dict[str, list[tuple[Any, ...]]]:
            resp = await _call(session, seeded.finance, "GET", _export("work-progress") + query)
            assert resp.status_code == 200, resp.text
            wb = load_workbook(io.BytesIO(resp.content), data_only=False)
            assert wb.sheetnames == ["PR 进度", "单品汇总", "周明细"]
            return {n: list(wb[n].iter_rows(values_only=True)) for n in wb.sheetnames}

        bare = await sheets("")
        autumn = await sheets("&season=2026秋")
        header = bare["PR 进度"][0]
        assert "推广成本" in header and "成本" not in header
        quote_col = header.index("约篇量")
        # 场景有效性：不带季节时 pr / pr2 / 未分配三行共 5 单；单品 = 季节商品 + 末行未归属
        assert sum(r[quote_col] for r in bare["PR 进度"][1:]) == 5
        assert [r[0] for r in bare["单品汇总"][1:]] == ["季节导出商品", "未归属商品"]
        assert sum(r[quote_col + 2] for r in bare["周明细"][1:]) == 5
        # 带季节：三个 sheet 都只剩命中的那一单，没有未归属行
        amount_col = header.index("约稿服务费")
        assert [(r[quote_col], r[amount_col]) for r in autumn["PR 进度"][1:]] == [(1, 66.0)]
        assert [(r[0], r[quote_col]) for r in autumn["单品汇总"][1:]] == [("季节导出商品", 1)]
        assert sum(r[quote_col + 2] for r in autumn["周明细"][1:]) == 1
        cells = {c for s in bare.values() for r in s for c in r}
        assert goods.goods_code not in cells

    @pytest.mark.parametrize("kind", ["work-progress", "production", "store-daily"])
    async def test_operations_all_types(
        self, session: AsyncSession, flow_users: Any, kind: str
    ) -> None:
        resp = await _call(session, flow_users.operations, "GET", _export(kind))
        assert resp.status_code == 200, resp.text

    async def test_follows_read_permission(self, session: AsyncSession, flow_users: Any) -> None:
        """判的是对应读权限本身：财务个人授予投产读后，导出投产放行，店铺照旧 403。"""
        fin = flow_users.finance
        before = await _call(session, fin, "GET", _export("production"))
        assert (before.status_code, before.json()["code"]) == (403, "PERMISSION_DENIED")
        assert before.json()["details"]["required_scope"] == "report.production"
        await _override(session, fin, "report.production:read", "grant")
        assert (await _call(session, fin, "GET", _export("production"))).status_code == 200
        store = await _call(session, fin, "GET", _export("store-daily"))
        assert store.json()["details"]["required_scope"] == "report.store_daily"

    async def test_invalid_type_still_400(self, session: AsyncSession, flow_users: Any) -> None:
        """类型先校验：不认识的类型照旧 400，不因为没有对应读权限变成 403。"""
        resp = await _call(session, flow_users.finance, "GET", _export("unknown"))
        assert (resp.status_code, resp.json()["code"]) == (400, "REPORT_EXPORT_TYPE_INVALID")

    @pytest.mark.parametrize("role", ["pr", "warehouse"])
    async def test_without_export_scope(
        self, session: AsyncSession, flow_users: Any, role: str
    ) -> None:
        resp = await _call(session, getattr(flow_users, role), "GET", _export("work-progress"))
        assert (resp.status_code, resp.json()["code"]) == (403, "PERMISSION_DENIED")
        assert resp.json()["details"]["required_scope"] == "report.export"

    async def test_read_own_with_export_grant_still_forbidden(
        self, session: AsyncSession, flow_users: Any
    ) -> None:
        """pr 个人授予 ``report.export:read``、只持 ``work_progress:read_own``：导出工作进度仍 403。"""
        pr = flow_users.pr
        # 场景有效性：pr 能看工作进度（只持 read_own）
        assert (await _call(session, pr, "GET", f"{_WP}?month={_MONTH}")).status_code == 200
        await _override(session, pr, "report.export:read", "grant")
        resp = await _call(session, pr, "GET", _export("work-progress"))
        assert (resp.status_code, resp.json()["code"]) == (403, "PERMISSION_DENIED")
        # 过了路由的 report.export 关，挡在 service 的工作进度 read 上
        details = resp.json()["details"]
        assert (details["required_scope"], details["required_action"]) == (
            "report.work_progress",
            "read",
        )


_CAPS = "/api/reports/capabilities"
_ALL_EXPORTS = ["work-progress", "production", "store-daily"]

# 6 个真实角色逐字段期望（报表视图设计第 4 节「能力端点」）
_CAPS_EXPECTED: dict[str, dict[str, Any]] = {
    "pr": {
        "work_progress": "own",
        "publish_progress": "own",
        "bi": "personal",
        "production": False,
        "production_margin": False,
        "store_read": False,
        "store_write": False,
        "export": [],
        "refresh": False,
    },
    "pr_manager": {
        "work_progress": "all",
        "publish_progress": "all",
        "bi": "full",
        "production": True,
        "production_margin": False,
        "store_read": True,
        "store_write": False,
        "export": _ALL_EXPORTS,
        "refresh": True,
    },
    "admin": {
        "work_progress": "all",
        "publish_progress": "all",
        "bi": "full",
        "production": True,
        "production_margin": True,
        "store_read": True,
        "store_write": True,
        "export": _ALL_EXPORTS,
        "refresh": True,
    },
    "finance": {
        "work_progress": "all",
        "publish_progress": None,
        "bi": "full",
        "production": False,
        "production_margin": False,
        "store_read": False,
        "store_write": False,
        "export": ["work-progress"],
        "refresh": False,
    },
    "operations": {
        "work_progress": "all",
        "publish_progress": "all",
        "bi": "full",
        "production": True,
        "production_margin": False,
        "store_read": True,
        "store_write": True,
        "export": _ALL_EXPORTS,
        "refresh": True,
    },
    "warehouse": {
        "work_progress": None,
        "publish_progress": None,
        "bi": None,
        "production": False,
        "production_margin": False,
        "store_read": False,
        "store_write": False,
        "export": [],
        "refresh": False,
    },
}


@pytest.mark.usefixtures("tenant_ctx")
class TestCapabilities:
    @pytest.mark.parametrize("role", sorted(_CAPS_EXPECTED))
    async def test_real_roles(self, session: AsyncSession, flow_users: Any, role: str) -> None:
        resp = await _call(session, getattr(flow_users, role), "GET", _CAPS)
        assert resp.status_code == 200, resp.text
        assert resp.json() == _CAPS_EXPECTED[role]

    async def test_table_is_effective(self) -> None:
        """场景有效性：期望表里每个字段至少出现两种值（不是全 null / 全 False 的空转）。"""
        for field in _CAPS_EXPECTED["pr"]:
            values = {repr(caps[field]) for caps in _CAPS_EXPECTED.values()}
            assert len(values) >= 2, field
        assert {c["work_progress"] for c in _CAPS_EXPECTED.values()} == {"all", "own", None}
        assert {c["bi"] for c in _CAPS_EXPECTED.values()} == {"full", "personal", None}

    async def test_export_follows_read_permission(
        self, session: AsyncSession, flow_users: Any
    ) -> None:
        """导出能力 = 持 ``report.export:read`` 且有该报表读权限：财务授予投产读后多出 production。"""
        fin = flow_users.finance
        await _override(session, fin, "report.production:read", "grant")
        resp = await _call(session, fin, "GET", _CAPS)
        assert resp.status_code == 200, resp.text
        assert resp.json()["export"] == ["work-progress", "production"]
        assert resp.json()["production"] is True
        # 能看投产不等于能看毛利（只管理员）
        assert resp.json()["production_margin"] is False

    async def test_margin_needs_production_read(self) -> None:
        """毛利标记 = 投产读 且 F1 可见：超管字段上下文、但没有投产读 → false。"""
        from app.core.security.field_permissions import FieldPermissionContext
        from app.core.security.permissions import EffectivePermissions
        from app.modules.report.capabilities_api import build_capabilities

        superuser = FieldPermissionContext(frozenset(), frozenset(), frozenset(), is_superuser=True)
        without = EffectivePermissions(user_id="u", scopes=frozenset({"report.bi:read"}))
        with_prod = EffectivePermissions(user_id="u", scopes=frozenset({"report.production:read"}))
        assert build_capabilities(without, superuser).production_margin is False
        assert build_capabilities(with_prod, superuser).production_margin is True

    async def test_requires_login(self, session: AsyncSession) -> None:
        from app.core.db import get_session
        from app.main import app

        async def _session_override() -> AsyncIterator[AsyncSession]:
            yield session

        app.dependency_overrides[get_session] = _session_override
        try:
            async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
                resp = await c.get(_CAPS)
        finally:
            app.dependency_overrides.pop(get_session, None)
        assert resp.status_code == 401


# ----------------------- 单品投产毛利三列（7e R3，F1） ----------------------- #

_MARGIN_KEYS = ("goods_cost_total", "gross_profit", "gross_margin_rate")


def _margin(row: dict[str, Any]) -> tuple[Any, ...]:
    return tuple(row[k] for k in _MARGIN_KEYS)


def _by_goods(rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    return {r["goods_id"]: r for r in rows}


@pytest.fixture
async def margin_seeded(session: AsyncSession, flow_users: Any, product_factory: Any) -> Any:
    """A：启用成员成本 30.00 + 20.50、另有一个停用成员没有成本，本期与上期都有销售；
    B：两个启用成员里一个没有成本，本期有销售。"""
    from types import SimpleNamespace

    from app.modules.collect.models import QianniuDaily
    from app.modules.product.goods_models import GoodsMain, GoodsStyleItem
    from app.modules.product.platform_product_models import PlatformProduct

    tenant = flow_users.tenant

    async def goods(members: list[tuple[str | None, bool]]) -> Any:
        code = f"MG{uuid4().hex[:8]}"
        g = GoodsMain(
            tenant_id=tenant.id,
            goods_code=code,
            goods_title=f"{code} 商品",
            season="春夏",
            is_suit=len(members) > 1,
        )
        session.add(g)
        await session.flush()
        items = []
        styles = []
        for idx, (cost, active) in enumerate(members):
            style = await product_factory.style(style_code=f"MS{uuid4().hex[:8]}")
            styles.append(style)
            item = GoodsStyleItem(
                tenant_id=tenant.id,
                goods_main_id=g.id,
                style_id=style.id,
                sort_order=idx,
                single_goods_cost=Decimal(cost) if cost is not None else None,
                is_active=active,
            )
            session.add(item)
            items.append(item)
        pp = PlatformProduct(
            tenant_id=tenant.id,
            platform="千牛",
            platform_id=f"P{uuid4().hex[:10]}",
            style_id=styles[0].id,
            goods_main_id=g.id,
            channel="普通",
            is_active=True,
        )
        session.add(pp)
        await session.flush()
        return SimpleNamespace(goods=g, items=items, pp=pp)

    def qn(pp: Any, day: Any, pay: str, refund: str, orders: int) -> QianniuDaily:
        return QianniuDaily(
            tenant_id=tenant.id,
            platform_product_id=pp.id,
            platform_id_snapshot=pp.platform_id,
            date=day,
            visitors=100,
            pay_amount=Decimal(pay),
            pay_orders=orders,
            refund_amount=Decimal(refund),
        )

    a = await goods([("30.00", True), ("20.50", True), (None, False)])
    b = await goods([("40.00", True), (None, True)])
    session.add_all(
        [
            qn(a.pp, _M_LO + timedelta(days=2), "1000.00", "100.00", 6),
            # 上一期（环比区间）
            qn(a.pp, _M_LO - timedelta(days=3), "200.00", "0.00", 2),
            qn(b.pp, _M_LO + timedelta(days=4), "500.00", "0.00", 3),
        ]
    )
    await session.flush()
    return SimpleNamespace(users=flow_users, a=a, b=b)


@pytest.mark.usefixtures("tenant_ctx")
class TestProductionMargin:
    """毛利三列只给管理员（10-10 业务答复）；成本读取时现算，任一启用成员缺成本 → 三列 null。"""

    async def test_admin_hand_computed(self, session: AsyncSession, margin_seeded: Any) -> None:
        resp = await _call(session, margin_seeded.users.admin, "GET", _PROD)
        assert resp.status_code == 200, resp.text
        items = _by_goods(resp.json()["items"])
        a, b = str(margin_seeded.a.goods.id), str(margin_seeded.b.goods.id)
        # A：实际销售 900.00；总成本 6 × (30.00 + 20.50) = 303.00（停用成员缺成本不影响）
        assert _margin(items[a]) == ("303.00", "597.00", "0.6633")
        # B：一个启用成员没有成本 → 三列都 null，不当 0
        assert _margin(items[b]) == (None, None, None)
        # 上期同规则：200.00 − 2 × 50.50
        assert _margin(_by_goods(resp.json()["previous"])[a]) == ("101.00", "99.00", "0.4950")

    @pytest.mark.parametrize("role", ["pr_manager", "operations"])
    async def test_non_admin_null_rest_equal(
        self, session: AsyncSession, margin_seeded: Any, role: str
    ) -> None:
        """运营能看成本价、主管看得到销售额：毛利三列都给 null，其余列与管理员逐字相同。"""
        users = margin_seeded.users
        admin = await _call(session, users.admin, "GET", _PROD)
        other = await _call(session, getattr(users, role), "GET", _PROD)
        assert admin.status_code == other.status_code == 200, other.text
        for part in ("items", "previous"):
            rows = other.json()[part]
            assert rows, part
            assert all(_margin(r) == (None, None, None) for r in rows), part

        def strip(body: dict[str, Any]) -> dict[str, Any]:
            return {
                k: [{c: v for c, v in r.items() if c not in _MARGIN_KEYS} for r in body[k]]
                if k in ("items", "previous")
                else body[k]
                for k in body
            }

        assert strip(other.json()) == strip(admin.json())
        # 场景有效性：管理员那边确实有数
        assert any(_margin(r)[0] is not None for r in admin.json()["items"])

    @pytest.mark.parametrize(
        ("role", "expected"),
        [
            ("admin", (303.0, 597.0, 0.6633)),
            ("pr_manager", (None, None, None)),
            ("operations", (None, None, None)),
        ],
    )
    async def test_export_same_rule(
        self, session: AsyncSession, margin_seeded: Any, role: str, expected: tuple[Any, ...]
    ) -> None:
        """导出同页面规则：三列表头都在；非管理员单元格为空，管理员有数。列序 = 共用列定义。"""
        import io

        from openpyxl import load_workbook

        from app.modules.report.production_columns import PRODUCTION_COLUMNS

        user = getattr(margin_seeded.users, role)
        resp = await _call(session, user, "GET", _export("production"))
        assert resp.status_code == 200, resp.text
        sheet = load_workbook(io.BytesIO(resp.content)).active
        rows = list(sheet.iter_rows(values_only=True))
        header = list(rows[0])
        fixed = header[: len(PRODUCTION_COLUMNS)]
        assert fixed[:4] == ["商品编码", "商品简称", "商品名称", "渠道"]
        recent7 = [h for h in fixed if h.startswith("近 7 天")]
        assert len(recent7) == 3 and all(h.startswith("近 7 天（") for h in recent7), recent7
        margin_idx = [header.index(h) for h in ("销售货品总成本", "单品毛利", "单品利润率")]
        by_code = {r[0]: r for r in rows[1:]}
        a_row = by_code[margin_seeded.a.goods.goods_code]
        assert a_row[header.index("渠道")] == "普通"
        got = tuple(None if a_row[i] in (None, "") else a_row[i] for i in margin_idx)
        assert got == expected
        # 非毛利列与角色无关：A 的支付金额都是 1000
        assert a_row[header.index("支付金额")] == 1000.0

    async def test_cost_change_takes_effect(
        self, session: AsyncSession, margin_seeded: Any
    ) -> None:
        """成本不进汇总表、读取时现算：改成员成本、补上缺的成本，下一次请求就变。"""
        a, b = margin_seeded.a, margin_seeded.b
        a.items[0].single_goods_cost = Decimal("10.00")
        b.items[1].single_goods_cost = Decimal("5.00")
        await session.flush()
        resp = await _call(session, margin_seeded.users.admin, "GET", _PROD)
        assert resp.status_code == 200, resp.text
        items = _by_goods(resp.json()["items"])
        # A：6 × 30.50 = 183.00；B：3 × 45.00 = 135.00
        assert _margin(items[str(a.goods.id)]) == ("183.00", "717.00", "0.7967")
        assert _margin(items[str(b.goods.id)]) == ("135.00", "365.00", "0.7300")


# ----------------------- 店铺日 / 周 / 月（7e R4，设计 2.4） ----------------------- #

_STORE_PAGES = ("daily", "weekly", "monthly")
_STORE_LABELS = {"daily": "store_daily", "weekly": "store_weekly", "monthly": "store_monthly"}
# 旧接口 /store-daily 的行键（保留一个版本，结构不变）
_LEGACY_STORE_KEYS = {
    "date",
    "visitors",
    "pay_amount",
    "pay_orders",
    "ad_spend_total",
    "zhitongche_spend",
    "yinli_spend",
    "extra",
}
# 上月第一个不在 1 号的周三（区间从周三起、不是整月）
_AFTER_FIRST = _M_LO + timedelta(days=1)
_WED = _AFTER_FIRST + timedelta(days=(2 - _AFTER_FIRST.weekday()) % 7)


def _store(page: str, lo: Any, hi: Any) -> str:
    return f"/api/reports/store/{page}?preset=custom&date_from={lo}&date_to={hi}"


@pytest.fixture
def store_recorder(monkeypatch: pytest.MonkeyPatch) -> list[tuple[str, str]]:
    """拦截店铺 service 的 record_source，记下每次读取走了哪条路径。"""
    import app.modules.report.store_daily_service as store_mod

    calls: list[tuple[str, str]] = []
    monkeypatch.setattr(store_mod, "record_source", lambda r, s: calls.append((r, s)))
    return calls


@pytest.mark.usefixtures("tenant_ctx")
class TestStoreReport:
    @pytest.mark.parametrize("page", _STORE_PAGES)
    async def test_finance_forbidden(
        self, session: AsyncSession, flow_users: Any, page: str
    ) -> None:
        """财务没有店铺读权限：三页都 403（菜单也不该出，见前端）。"""
        resp = await _call(session, flow_users.finance, "GET", _store(page, _M_LO, _M_HI))
        assert (resp.status_code, resp.json()["code"]) == (403, "PERMISSION_DENIED")
        assert resp.json()["details"]["required_scope"] == "report.store_daily"

    @pytest.mark.parametrize("page", _STORE_PAGES)
    async def test_pr_manager_reads(
        self, session: AsyncSession, flow_users: Any, page: str
    ) -> None:
        resp = await _call(session, flow_users.pr_manager, "GET", _store(page, _M_LO, _M_HI))
        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["granularity"] == {"daily": "day", "weekly": "week", "monthly": "month"}[page]
        assert (body["date_from"], body["date_to"]) == (str(_M_LO), str(_M_HI))
        assert set(body) == {
            "granularity",
            "date_from",
            "date_to",
            "items",
            "recorded_days",
            "total_days",
        }

    @pytest.mark.parametrize("page", _STORE_PAGES)
    @pytest.mark.parametrize(
        ("query", "code"),
        [
            ("preset=custom&date_from=2025-01-01&date_to=2026-01-03", "REPORT_INVALID_TIME_RANGE"),
            ("preset=custom&date_from=2026-03-02&date_to=2026-03-01", "REPORT_INVALID_TIME_RANGE"),
            ("preset=last_year", "REPORT_INVALID_TIME_PRESET"),
        ],
        ids=["span_over_366", "from_after_to", "unknown_preset"],
    )
    async def test_invalid_range_422(
        self, session: AsyncSession, flow_users: Any, page: str, query: str, code: str
    ) -> None:
        resp = await _call(
            session, flow_users.pr_manager, "GET", f"/api/reports/store/{page}?{query}"
        )
        assert (resp.status_code, resp.json()["code"]) == (422, code)

    async def test_new_rows_real_bounds_legacy_unchanged(
        self, session: AsyncSession, flow_users: Any
    ) -> None:
        """周三手填一项：新周表行上是实际起（周三）；旧接口仍是原来 8 个键、date = 桶首日（周一）。"""
        ops = flow_users.operations
        put = await _call(
            session, ops, "PUT", f"/api/reports/store-daily/{_WED}", {"zhitongche_spend": "40.00"}
        )
        assert put.status_code == 200, put.text
        hi = _WED + timedelta(days=6)

        weekly = await _call(session, ops, "GET", _store("weekly", _WED, hi))
        assert weekly.status_code == 200, weekly.text
        (row,) = weekly.json()["items"]
        assert (row["date_from"], row["date_to"]) == (str(_WED), str(_WED + timedelta(days=4)))
        assert (row["zhitongche_spend"], row["internal_spend"]) == ("40.00", "40.00")
        assert weekly.json()["total_days"] == 1

        legacy_q = f"preset=custom&date_from={_WED}&date_to={hi}"
        for granularity, first_date in (("day", _WED), ("week", _WED - timedelta(days=2))):
            legacy = await _call(
                session,
                ops,
                "GET",
                f"/api/reports/store-daily?{legacy_q}&granularity={granularity}",
            )
            assert legacy.status_code == 200, legacy.text
            rows = legacy.json()
            assert rows, granularity
            assert all(set(r) == _LEGACY_STORE_KEYS for r in rows), rows[0]
            assert rows[0]["date"] == str(first_date)
            assert rows[0]["zhitongche_spend"] == "40.00"

    async def test_one_source_per_response(
        self,
        session: AsyncSession,
        flow_users: Any,
        store_recorder: list[tuple[str, str]],
    ) -> None:
        """整段只判一次来源：只覆盖了前 7 天时，跨两周以上的请求整段实时、只记一次。"""
        from app.modules.report.summary_refresh_service import SummaryRefreshService

        await SummaryRefreshService(session).refresh(
            tenant_id=flow_users.tenant.id, date_from=_M_LO, date_to=_M_LO + timedelta(days=6)
        )
        await session.commit()
        mgr = flow_users.pr_manager
        for page in _STORE_PAGES:
            store_recorder.clear()
            covered = await _call(session, mgr, "GET", _store(page, _M_LO, _M_LO + timedelta(6)))
            assert covered.status_code == 200, covered.text
            # 场景有效性：被覆盖的区间真的走汇总
            assert store_recorder == [(_STORE_LABELS[page], "summary")], page
            store_recorder.clear()
            partial = await _call(session, mgr, "GET", _store(page, _M_LO, _M_LO + timedelta(20)))
            assert partial.status_code == 200, partial.text
            assert store_recorder == [(_STORE_LABELS[page], "live")], page


_PUT_STORE = "/api/reports/store-daily"
_MANUAL_NONE = {
    "ad_spend_total": None,
    "zhitongche_spend": None,
    "yinli_spend": None,
    "refund_only_amount": None,
    "remark": None,
}


@pytest.mark.usefixtures("tenant_ctx")
class TestStoreDailyPut:
    """``PUT /store-daily/{day}``：仅退款、校验、``null`` 清空（设计 2.4、5.1）。"""

    async def test_future_date_rejected_boundaries_ok(
        self, session: AsyncSession, flow_users: Any
    ) -> None:
        """明天 → 422 且不落库；今天、2020-01-01、上限金额、200 字备注都能存。"""
        ops = flow_users.operations
        tomorrow = _TODAY + timedelta(days=1)
        resp = await _call(
            session, ops, "PUT", f"{_PUT_STORE}/{tomorrow}", {"refund_only_amount": "1.00"}
        )
        assert (resp.status_code, resp.json()["code"]) == (422, "STORE_DAILY_FUTURE_DATE")
        written = await session.execute(
            text(
                "SELECT count(*) FROM store_daily "
                "WHERE tenant_id = CAST(:t AS uuid) AND date = CAST(:d AS date)"
            ),
            {"t": str(flow_users.tenant.id), "d": tomorrow},
        )
        assert written.scalar_one() == 0

        body = {"refund_only_amount": "99999999.99", "remark": "备" * 200}
        for day in (_TODAY, "2020-01-01"):
            ok = await _call(session, ops, "PUT", f"{_PUT_STORE}/{day}", body)
            assert ok.status_code == 200, ok.text
            assert ok.json() == {"ok": True, "date": str(day), **_MANUAL_NONE, **body}

    @pytest.mark.parametrize(
        ("day", "body", "code"),
        [
            ("2019-12-31", {"refund_only_amount": "1.00"}, "STORE_DAILY_DATE_TOO_EARLY"),
            (None, {"zhitongche_spend": "-0.01"}, "VALIDATION_ERROR"),
            (None, {"yinli_spend": "100000000.00"}, "VALIDATION_ERROR"),
            (None, {"ad_spend_total": "1.234"}, "VALIDATION_ERROR"),
            (None, {"refund_only_amount": "abc"}, "VALIDATION_ERROR"),
            (None, {"remark": "备" * 201}, "VALIDATION_ERROR"),
            (None, {}, "VALIDATION_ERROR"),
        ],
        ids=[
            "before_2020",
            "negative",
            "over_max",
            "three_decimals",
            "not_number",
            "remark_201",
            "empty",
        ],
    )
    async def test_invalid_422(
        self, session: AsyncSession, flow_users: Any, day: str | None, body: Any, code: str
    ) -> None:
        resp = await _call(
            session, flow_users.operations, "PUT", f"{_PUT_STORE}/{day or _M_LO}", body
        )
        assert (resp.status_code, resp.json()["code"]) == (422, code), resp.text

    async def test_pr_manager_forbidden(self, session: AsyncSession, flow_users: Any) -> None:
        """主管只读：PUT 403。"""
        resp = await _call(
            session,
            flow_users.pr_manager,
            "PUT",
            f"{_PUT_STORE}/{_M_LO}",
            {"refund_only_amount": "1.00"},
        )
        assert (resp.status_code, resp.json()["code"]) == (403, "PERMISSION_DENIED")
        assert resp.json()["details"]["required_scope"] == "report.store_daily"

    async def test_null_clears_via_api(self, session: AsyncSession, flow_users: Any) -> None:
        """传 ``null`` 清空、没传的不动；日表立即读到。"""
        ops = flow_users.operations
        first = await _call(
            session,
            ops,
            "PUT",
            f"{_PUT_STORE}/{_WED}",
            {"zhitongche_spend": "40.00", "refund_only_amount": "5.00"},
        )
        assert first.status_code == 200, first.text
        cleared = await _call(
            session, ops, "PUT", f"{_PUT_STORE}/{_WED}", {"zhitongche_spend": None}
        )
        assert cleared.status_code == 200, cleared.text
        assert cleared.json() == {
            "ok": True,
            "date": str(_WED),
            **_MANUAL_NONE,
            "refund_only_amount": "5.00",
        }
        daily = await _call(session, ops, "GET", _store("daily", _WED, _WED))
        assert daily.status_code == 200, daily.text
        (row,) = daily.json()["items"]
        assert (row["zhitongche_spend"], row["refund_only_amount"]) == (None, "5.00")
        assert (row["refund_only_days"], row["day_count"]) == (1, 1)
