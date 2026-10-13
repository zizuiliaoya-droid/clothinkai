"""BI / 报表导出 API 契约测试。"""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal
from types import SimpleNamespace
from uuid import uuid4

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook

from app.core.security.field_permissions import FieldPermissionContext
from app.core.security.permissions import EffectivePermissions
from app.modules.report.advanced_schemas import ProductionReport, ProductionRow, Recent7Range
from app.modules.report.pr_scope import PR_SCOPE_ALL
from app.modules.report.production_columns import PRODUCTION_COLUMNS

_SUPER_CTX = FieldPermissionContext(frozenset(), frozenset(), frozenset(), is_superuser=True)
_ALL_PERMS = EffectivePermissions(user_id="test", scopes=frozenset({"*"}))


@pytest.mark.api
@pytest.mark.asyncio
class TestReportExportApiContract:
    async def test_bi_dashboard_requires_auth(self) -> None:
        from app.main import app

        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
            resp = await ac.get("/api/reports/bi")
        assert resp.status_code == 401

    async def test_export_requires_auth(self) -> None:
        from app.main import app

        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
            resp = await ac.get("/api/reports/production/export")
        assert resp.status_code == 401

    async def test_openapi_exposes_bi_and_export(self) -> None:
        from app.main import app

        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
            resp = await ac.get("/api/openapi.json")
        assert resp.status_code == 200
        paths = resp.json().get("paths", {})
        assert "/api/reports/bi" in paths
        assert "/api/reports/bi/layout" in paths
        assert "/api/reports/{report_type}/export" in paths

    async def test_export_rows_are_complete_and_numeric(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        from app.modules.report import export_service as module

        dangerous_header = '=HYPERLINK("https://invalid.example","x")'
        dangerous_text = "  +SUM(1,1)"
        dangerous_short = "=1+1"
        goods_id = uuid4()
        production_row = ProductionRow(
            goods_id=goods_id,
            goods_code="S001",
            goods_short_name=dangerous_short,
            goods_title=dangerous_text,
            style_codes=["S001", "S002"],
            channel="普通",
            link_count=2,
            pay_amount=Decimal("100"),
            refund_amount=Decimal("10"),
            return_rate=Decimal("0.1"),
            confirmed_amount=Decimal("90"),
            promo_cost=Decimal("20"),
            ad_spend=Decimal("30"),
            total_spend=Decimal("50"),
            pay_orders=4,
            add_cart_count=5,
            add_cart_cost=Decimal("10"),
            ad_add_cart_count=3,
            ad_add_cart_cost=Decimal("10"),
            offsite_add_cart=2,
            pre_return_roi=Decimal("2"),
            net_roi=Decimal("1.8"),
            unit_deal_cost=Decimal("12.5"),
            recent7_pay_amount=Decimal("70"),
            recent7_total_spend=Decimal("7"),
            recent7_net_roi=Decimal("10"),
            goods_cost_total=Decimal("40"),
            gross_profit=Decimal("50"),
            gross_margin_rate=Decimal("0.5556"),
            extra={dangerous_header: "3"},
        )
        # 同一商品的直播行：extra 只挂在普通行（S1 偏差），直播行这一列留空
        live_row = production_row.model_copy(
            update={"channel": "直播", "link_count": 1, "offsite_add_cart": None, "extra": {}}
        )
        production_report = ProductionReport(
            items=[production_row, live_row],
            recent7_range=Recent7Range(date_from=date(2026, 7, 27), date_to=date(2026, 8, 2)),
        )
        production_calls: list[dict[str, object]] = []
        # service 给什么行导出就写什么行：分桶与 extra 聚合都在 StoreDailyService 里。
        # 两行故意落在同一周 —— 导出要是自己再按周合并一次，这里会变成一行。
        store_rows = [
            SimpleNamespace(
                date=date(2026, 8, 3),
                visitors=10,
                pay_amount=Decimal("20"),
                pay_orders=2,
                ad_spend_total=Decimal("3"),
                zhitongche_spend=Decimal("1"),
                yinli_spend=None,
                # 多行时算不出来的比率是 None → 导出为空单元格
                extra={"成交人数": "2", "下单转化率": None},
            ),
            SimpleNamespace(
                date=date(2026, 8, 4),
                visitors=5,
                pay_amount=Decimal("8"),
                pay_orders=1,
                ad_spend_total=None,
                zhitongche_spend=Decimal("2"),
                yinli_spend=Decimal("4"),
                extra={"成交人数": "1", "下单转化率": "0.63"},
            ),
        ]
        store_calls: list[dict[str, object]] = []
        work_calls: list[tuple[tuple[object, ...], dict[str, object]]] = []
        work_row = SimpleNamespace(
            pr_name="PR甲",
            quote_count=10,
            in_schedule_count=1,
            urge_count=2,
            important_urge_count=3,
            overdue_count=4,
            publish_count=5,
            info_complete_count=4,
            info_complete_rate=Decimal("0.8"),
            cancel_count=1,
            recall_due_count=2,
            recall_success_count=1,
            recall_complete_rate=Decimal("0.5"),
            overdue_rate=Decimal("0.4"),
            month_complete_rate=Decimal("0.5"),
            hit_count=2,
            hit_rate=Decimal("0.4"),
            like_count=100,
            cost=Decimal("200"),
            cpl=Decimal("2"),
        )

        class FakeProductionService:
            def __init__(self, _session: object) -> None:
                pass

            async def get_report(self, *_args: object, **kwargs: object) -> object:
                production_calls.append(kwargs)
                return production_report

        class FakeStoreDailyService:
            def __init__(self, _session: object) -> None:
                pass

            async def get_dashboard(self, *_args: object, **kwargs: object) -> list[object]:
                store_calls.append(kwargs)
                return store_rows

        class FakeWorkProgressService:
            def __init__(self, _session: object) -> None:
                pass

            async def get_for_range(self, *args: object, **kwargs: object) -> list[object]:
                work_calls.append((args[1:], kwargs))
                return [work_row]

        monkeypatch.setattr(module, "ProductionService", FakeProductionService)
        monkeypatch.setattr(module, "StoreDailyService", FakeStoreDailyService)
        monkeypatch.setattr(module, "WorkProgressService", FakeWorkProgressService)
        service = module.ReportExportService(None)  # type: ignore[arg-type]
        period = (date(2026, 8, 3), date(2026, 8, 4))

        production_headers, production_rows = await service._fetch_rows(
            uuid4(),
            "production",
            period,
            exclude_brushing=True,
            seasons=None,
            granularity="day",
            field_ctx=_SUPER_CTX,
        )
        # 表头逐列来自共用列定义（列顺序 = 页面），近 7 天三列带区间；extra 在最后
        assert production_headers == [
            "商品编码",
            "商品简称",
            "商品名称",
            "渠道",
            "链接数",
            "含款号",
            "支付金额",
            "支付件数（不扣刷单）",
            "退款金额",
            "退货退款率",
            "实际销售金额",
            "站外推广费",
            "站内投放",
            "推广总花费",
            "退货前投产比",
            "净投产比",
            "推广单件成交成本",
            "总加购数",
            "总加购成本",
            "站内加购",
            "站内加购成本",
            "站外加购",
            "近 7 天（07-27 ~ 08-02）支付金额",
            "近 7 天（07-27 ~ 08-02）推广总花费",
            "近 7 天（07-27 ~ 08-02）净投产比",
            "销售货品总成本",
            "单品毛利",
            "单品利润率",
            dangerous_header,
        ]
        assert production_headers[:-1] == [
            module.export_header(c, production_report.recent7_range) for c in PRODUCTION_COLUMNS
        ]
        assert production_rows[0][:-1] == [
            "S001",
            dangerous_short,
            dangerous_text,
            "普通",
            2,
            "S001,S002",
            Decimal("100"),
            4,
            Decimal("10"),
            Decimal("0.1"),
            Decimal("90"),
            Decimal("20"),
            Decimal("30"),
            Decimal("50"),
            Decimal("2"),
            Decimal("1.8"),
            Decimal("12.5"),
            5,
            Decimal("10"),
            3,
            Decimal("10"),
            2,
            Decimal("70"),
            Decimal("7"),
            Decimal("10"),
            Decimal("40"),
            Decimal("50"),
            Decimal("0.5556"),
        ]
        # 同商品两渠道 → 两行；直播行站外加购 None、extra 空
        assert [r[3] for r in production_rows] == ["普通", "直播"]
        live = dict(zip(production_headers, production_rows[1], strict=True))
        assert live["站外加购"] is None and live[dangerous_header] is None
        assert len(production_headers) == len(production_rows[0]) == len(PRODUCTION_COLUMNS) + 1
        assert isinstance(module._cell(production_rows[0][6]), float)
        assert isinstance(module._cell(production_rows[0][-1]), float)
        # 超管 ctx → 毛利可见，原样交给 service；运营 ctx（能看成本价）→ 不给毛利
        assert production_calls[-1]["show_margin"] is True
        ops_ctx = FieldPermissionContext(frozenset({"operations"}), frozenset(), frozenset())
        await service._fetch_rows(
            uuid4(),
            "production",
            period,
            exclude_brushing=True,
            seasons=None,
            granularity="day",
            field_ctx=ops_ctx,
        )
        assert production_calls[-1]["show_margin"] is False
        assert module._cell(dangerous_header) == f"'{dangerous_header}"
        assert module._cell(dangerous_text) == f"'{dangerous_text}"
        assert module._cell(dangerous_short) == f"'{dangerous_short}"

        store_headers, grouped_store_rows = await service._fetch_rows(
            uuid4(),
            "store-daily",
            period,
            exclude_brushing=True,
            seasons=None,
            granularity="week",
            field_ctx=_SUPER_CTX,
        )
        assert store_headers[:7] == [
            "日期",
            "访客数",
            "支付额",
            "支付订单",
            "广告花费",
            "直通车花费",
            "引力魔方花费",
        ]
        # 粒度原样交给 service，导出不再自己分桶、相加
        assert store_calls == [{"granularity": "week"}]
        assert store_headers[7:] == ["下单转化率", "成交人数"]
        assert grouped_store_rows == [
            [
                date(2026, 8, 3),
                10,
                Decimal("20"),
                2,
                Decimal("3"),
                Decimal("1"),
                None,
                None,
                Decimal("2"),
            ],
            [
                date(2026, 8, 4),
                5,
                Decimal("8"),
                1,
                None,
                Decimal("2"),
                Decimal("4"),
                Decimal("0.63"),
                Decimal("1"),
            ],
        ]

        work_headers, work_rows = await service._fetch_rows(
            uuid4(),
            "work-progress",
            period,
            exclude_brushing=True,
            seasons=None,
            granularity="day",
            field_ctx=_SUPER_CTX,
        )
        assert len(work_headers) == len(work_rows[0]) == 20
        assert work_headers[1:9] == [
            "约篇量",
            "档期内",
            "催发",
            "重要催发",
            "超时",
            "发布量",
            "信息完整数",
            "信息完整率",
        ]
        assert work_rows[0][7:9] == [4, Decimal("0.8")]
        # 工作进度按请求区间导（不再取起日所在整月），看全部 PR
        assert work_calls == [(period, {"scope": PR_SCOPE_ALL, "field_ctx": _SUPER_CTX})]

        response = await service.export(
            uuid4(),
            "production",
            period,
            granularity="day",
            field_ctx=_SUPER_CTX,
            perms=_ALL_PERMS,
        )
        body = b"".join([chunk async for chunk in response.body_iterator])
        worksheet = load_workbook(io.BytesIO(body), data_only=False).active
        # 危险表头是唯一的动态 extra 列，落在固定列之后（最后一列）
        last_col = worksheet.max_column
        assert worksheet.cell(1, last_col).value == f"'{dangerous_header}"
        assert worksheet.cell(1, last_col).data_type != "f"
        # 第 2 列简称、第 3 列名称，两列都是用户可编辑的文字，都不能被当成公式执行
        assert worksheet.cell(2, 2).value == f"'{dangerous_short}"
        assert worksheet.cell(2, 2).data_type != "f"
        assert worksheet.cell(2, 3).value == f"'{dangerous_text}"
        assert worksheet.cell(2, 3).data_type != "f"

        store_response = await service.export(
            uuid4(),
            "store-daily",
            period,
            granularity="week",
            field_ctx=_SUPER_CTX,
            perms=_ALL_PERMS,
        )
        store_body = b"".join([chunk async for chunk in store_response.body_iterator])
        store_sheet = load_workbook(io.BytesIO(store_body), data_only=False).active
        assert store_calls[-1] == {"granularity": "week"}
        assert store_sheet.max_row == 3  # 表头 + service 给的两行，没有再合并
        # 第 8 列「下单转化率」：None 写成空单元格，有值的是 Excel 数值
        assert store_sheet.cell(1, 8).value == "下单转化率"
        assert store_sheet.cell(2, 8).value in (None, "")
        assert store_sheet.cell(3, 8).value == 0.63
