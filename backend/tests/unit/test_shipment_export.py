"""仓库待打单导出 ``promotion/shipment_export.py``（流程线 7.4、9.3 A「导出」）。

一行一个商品明细（套装 2 行、收件信息重复）；没有明细的旧单一行、颜色列放 ``legacy_color_spec`` 原文；
编码列保留；``watermark=None`` 与不传时输出字节相同。纯函数，不碰库。
"""

from __future__ import annotations

import io
from datetime import UTC, datetime

import pytest
from freezegun import freeze_time
from openpyxl import load_workbook

from app.modules.promotion.shipment_export import (
    SHIPMENT_EXPORT_HEADERS,
    ExportWatermark,
    ShipmentExportItem,
    ShipmentExportRow,
    build_shipment_workbook,
)

_PUSHED = datetime(2026, 10, 9, 2, 5, tzinfo=UTC)  # 北京时间 10:05


def _suit_row() -> ShipmentExportRow:
    return ShipmentExportRow(
        internal_code="DE2610090001",
        ship_pushed_at=_PUSHED,
        receiver_name="张三",
        receiver_phone="13812345678",
        receiver_address="杭州某路 1 号",
        style_code="ST-TOP",
        goods_code="G-SUIT",
        sku_code=None,
        short_name="条纹套装",
        legacy_color_spec=None,
        items=(
            ShipmentExportItem(
                style_code="ST-TOP", sku_code="SK-TOP-M", short_name="上衣", color="黑色", size="M"
            ),
            ShipmentExportItem(
                style_code="ST-PANTS",
                sku_code="SK-PANTS-L",
                short_name="阔腿裤",
                color="白色",
                size="L",
            ),
        ),
    )


def _legacy_row() -> ShipmentExportRow:
    return ShipmentExportRow(
        internal_code="DE2610090002",
        ship_pushed_at=None,
        receiver_name=None,
        receiver_phone=None,
        receiver_address="=HYPERLINK(1)",
        style_code="ST-OLD",
        goods_code=None,
        sku_code="SK-OLD",
        short_name="旧款",
        legacy_color_spec="黑色 M 码",
        items=(),
    )


def _sheet_values(data: bytes) -> list[tuple[object, ...]]:
    ws = load_workbook(io.BytesIO(data), read_only=True).worksheets[0]
    return [tuple(row) for row in ws.iter_rows(values_only=True)]


def test_headers_exact() -> None:
    assert SHIPMENT_EXPORT_HEADERS == (
        "内部编码",
        "推送时间",
        "收件人",
        "电话",
        "地址",
        "款式编码",
        "商品编码",
        "SKU 编码",
        "商品简称",
        "颜色",
        "尺码",
        "数量",
    )


def test_suit_two_lines_and_legacy_one_line() -> None:
    rows = _sheet_values(build_shipment_workbook([_suit_row(), _legacy_row()]))
    assert rows[0] == SHIPMENT_EXPORT_HEADERS
    assert rows[1:] == [
        (
            "DE2610090001",
            "2026-10-09 10:05",
            "张三",
            "13812345678",
            "杭州某路 1 号",
            "ST-TOP",
            "G-SUIT",
            "SK-TOP-M",
            "上衣",
            "黑色",
            "M",
            1,
        ),
        (
            "DE2610090001",
            "2026-10-09 10:05",
            "张三",
            "13812345678",
            "杭州某路 1 号",
            "ST-PANTS",
            "G-SUIT",
            "SK-PANTS-L",
            "阔腿裤",
            "白色",
            "L",
            1,
        ),
        # 旧单：一行，颜色列放原文、尺码空；以 = 开头的文本转义成字面量（不当公式）
        (
            "DE2610090002",
            None,
            None,
            None,
            "'=HYPERLINK(1)",
            "ST-OLD",
            None,
            "SK-OLD",
            "旧款",
            "黑色 M 码",
            None,
            1,
        ),
    ]


def test_empty_is_header_only() -> None:
    assert _sheet_values(build_shipment_workbook([])) == [SHIPMENT_EXPORT_HEADERS]


def test_watermark_none_same_bytes_as_omitted() -> None:
    rows = [_suit_row(), _legacy_row()]
    with freeze_time("2026-10-09 10:00:00"):
        omitted = build_shipment_workbook(rows)
        explicit = build_shipment_workbook(rows, watermark=None)
    assert omitted == explicit


def test_watermark_not_implemented_yet() -> None:
    """水印随批次 6；本轮调用方一律传 None，传了值直接报错，免得以为已经打上了。"""
    mark = ExportWatermark(employee_no="W001", name="仓库甲", issued_at=_PUSHED)
    with pytest.raises(NotImplementedError):
        build_shipment_workbook([_suit_row()], watermark=mark)
