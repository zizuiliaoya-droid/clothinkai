"""8b §6.1：导入文件版式 XlsxLayout（指定 sheet、表头标记、分组行给重名列加前缀、required）。

无版式时 ``_parse_rows`` 必须与 8b 之前的读法逐字一致：下面的 ``_legacy_parse_xlsx`` 是改动前的原样副本。
样本用脱敏造数（假昵称、假 ID），表头照灰豚「抖音博主库」的两行表头造。
"""

from __future__ import annotations

import io
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from app.modules.importer.file_layout import (
    LayoutNotFoundError,
    XlsxLayout,
    rename_duplicate_headers,
)
from app.tasks.import_tasks import _parse_failure_reason, _parse_rows

DOUYIN = XlsxLayout(
    sheet="抖音博主库",
    header_marker="博主ID",
    group_prefix_duplicates=True,
    required=True,
    origin="灰豚",
)


def _legacy_parse_xlsx(raw: bytes) -> list[tuple[int, dict[str, Any]]]:
    """8b 之前 ``import_tasks._parse_xlsx`` 的原样副本（回归基准，不要改）。"""
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        ws = wb.active
        rows: list[tuple[int, dict[str, Any]]] = []
        header: list[str] = []
        row_number = 0
        for excel_row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() if c is not None else "" for c in excel_row]
            if not header:
                if all(c == "" for c in cells):
                    continue
                header = cells
                continue
            row_number += 1
            record = {
                (header[j] if j < len(header) and header[j] else f"col_{j}"): (
                    "" if cell is None else str(cell)
                )
                for j, cell in enumerate(excel_row)
            }
            rows.append((row_number, record))
        return rows
    finally:
        wb.close()


def _book(sheets: dict[str, list[list[Any]]], active: str | None = None) -> bytes:
    wb = Workbook()
    first = True
    for title, data in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = title
        for r, values in enumerate(data, start=1):
            for c, v in enumerate(values, start=1):
                if v is not None:
                    ws.cell(row=r, column=c, value=v)
    if active is not None:
        wb.active = wb.sheetnames.index(active)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# 灰豚「抖音博主库」样子：第 1 行分组（合并格只有左上有值）、第 2 行表头、重名列、表头后面一串空列
_GROUP = ["达人概况", None, None, None, "30天视频分析", None, None, "7天视频分析", None, None]
_HEADER = [
    "抖音博主",
    "博主ID",
    "网页ID",
    "粉丝总量",
    "点赞",
    "评论",
    "互动量",
    "点赞",
    "评论",
    "互动量",
    None,
    None,
]
_ROW1 = ["测试博主甲", "T00001", 90000000001, "1.2w", "3.5万", "--", 100, "1,234", "", 7]
_ROW2 = ["测试博主乙", 12345, None, "1亿", None, None, None, None, None, None, None, "多余"]


def _douyin_book(active: str = "说明") -> bytes:
    return _book(
        {
            "说明": [["这是说明页"]],
            "抖音博主库": [_GROUP, _HEADER, _ROW1, [], _ROW2],
            "7天视频详情": [["博主ID"]],
        },
        active=active,
    )


class TestNoLayoutUnchanged:
    @pytest.mark.parametrize(
        "data",
        [
            [["name", "qty"], ["杯子", 5], ["碗", 8]],
            # 前置空行 / 标题行、表头首尾空白、空表头格、重名列（后者覆盖前者）、多余列、空数据行、数字与 None
            [
                [],
                [],
                [" 标题 ", "", "数量", "数量", None],
                ["a", 1, 2, 3, None, "多余"],
                [],
                [None, 1.5, None, "x"],
            ],
            # 分组行在上：无版式时它就是表头
            [_GROUP, _HEADER, _ROW1, _ROW2],
            [],
        ],
    )
    def test_same_as_legacy(self, data: list[list[Any]]) -> None:
        raw = _book({"Sheet": data})
        legacy = _legacy_parse_xlsx(raw)
        assert _parse_rows(raw, "data.xlsx") == legacy
        assert _parse_rows(raw, "data.xlsx", None) == legacy
        # 什么都没声明的版式也一样
        assert _parse_rows(raw, "data.xlsx", XlsxLayout()) == legacy

    def test_reads_active_sheet_without_layout(self) -> None:
        raw = _douyin_book(active="说明")
        assert _parse_rows(raw, "x.xlsx") == _legacy_parse_xlsx(raw) == []

    def test_csv_ignores_layout(self) -> None:
        raw = "抖音博主,博主ID\n测试博主甲,T00001\n".encode()
        assert (
            _parse_rows(raw, "x.csv", DOUYIN)
            == _parse_rows(raw, "x.csv")
            == [(1, {"抖音博主": "测试博主甲", "博主ID": "T00001"})]
        )


class TestDouyinLayout:
    def test_sheet_marker_and_group_prefix(self) -> None:
        rows = _parse_rows(_douyin_book(), "抖音.xlsx", DOUYIN)
        assert [n for n, _ in rows] == [1, 2, 3]  # 从表头下一行起数，空行也占号
        first = rows[0][1]
        assert list(first) == [
            "抖音博主",
            "博主ID",
            "网页ID",
            "粉丝总量",
            "30天视频分析·点赞",
            "30天视频分析·评论",
            "30天视频分析·互动量",
            "7天视频分析·点赞",
            "7天视频分析·评论",
            "7天视频分析·互动量",
            "col_10",  # 表头空的列照旧叫 col_j（read_only 按 sheet 宽度补齐）
            "col_11",
        ]
        assert first["博主ID"] == "T00001"
        assert first["网页ID"] == "90000000001"
        assert first["30天视频分析·点赞"] == "3.5万"
        assert first["7天视频分析·点赞"] == "1,234"
        assert first["30天视频分析·评论"] == "--"
        third = rows[2][1]
        assert third["博主ID"] == "12345"
        assert third["网页ID"] == ""
        assert third["col_11"] == "多余"
        assert all(v == "" for v in rows[1][1].values())  # 中间的空行照样产出（与无版式一致）

    def test_marker_after_title_rows(self) -> None:
        raw = _book(
            {"抖音博主库": [["灰豚导出"], [], ["导出时间", "2026-10-01"], _GROUP, _HEADER, _ROW1]}
        )
        rows = _parse_rows(raw, "x.xlsx", DOUYIN)
        assert len(rows) == 1
        assert rows[0][0] == 1
        assert rows[0][1]["30天视频分析·点赞"] == "3.5万"

    def test_marker_only_in_first_ten_non_empty_rows(self) -> None:
        filler = [[f"说明{i}"] for i in range(10)]
        raw = _book({"抖音博主库": [*filler, _HEADER, _ROW1]})
        not_required = XlsxLayout(sheet="抖音博主库", header_marker="博主ID")
        rows = _parse_rows(raw, "x.xlsx", not_required)
        # 找不到 → 回落到第一个非空行当表头
        assert rows[0][0] == 1
        assert rows[0][1]["说明0"] == "说明1"
        with pytest.raises(LayoutNotFoundError):
            _parse_rows(raw, "x.xlsx", DOUYIN)

    def test_marker_must_equal_a_whole_cell(self) -> None:
        raw = _book({"抖音博主库": [["博主ID说明"], [" 博主ID ", "抖音博主"], ["T1", "甲"]]})
        rows = _parse_rows(raw, "x.xlsx", DOUYIN)
        assert rows == [(1, {"博主ID": "T1", "抖音博主": "甲"})]


class TestFallbackAndRequired:
    def test_missing_sheet_falls_back_to_active(self) -> None:
        raw = _book({"Sheet1": [["博主ID", "抖音博主"], ["T1", "甲"]]})
        layout = XlsxLayout(sheet="抖音博主库", header_marker="博主ID")
        assert _parse_rows(raw, "x.xlsx", layout) == [(1, {"博主ID": "T1", "抖音博主": "甲"})]

    def test_missing_marker_falls_back_to_first_non_empty(self) -> None:
        raw = _book({"抖音博主库": [[], ["账号", "昵称"], ["a1", "甲"]]})
        layout = XlsxLayout(sheet="抖音博主库", header_marker="博主ID")
        assert _parse_rows(raw, "x.xlsx", layout) == _legacy_parse_xlsx(raw)

    def test_required_missing_sheet(self) -> None:
        raw = _book({"Sheet1": [["博主ID", "抖音博主"], ["T1", "甲"]]})
        with pytest.raises(LayoutNotFoundError) as exc_info:
            _parse_rows(raw, "x.xlsx", DOUYIN)
        assert (
            str(exc_info.value)
            == "没有找到『抖音博主库』sheet 或『博主ID』表头，请上传灰豚导出的原文件"
        )

    def test_required_missing_marker(self) -> None:
        raw = _book({"抖音博主库": [["账号", "昵称"], ["a1", "甲"]]})
        with pytest.raises(LayoutNotFoundError):
            _parse_rows(raw, "x.xlsx", DOUYIN)

    def test_required_empty_sheet(self) -> None:
        raw = _book({"抖音博主库": []})
        with pytest.raises(LayoutNotFoundError):
            _parse_rows(raw, "x.xlsx", DOUYIN)

    def test_failure_reason(self) -> None:
        exc = LayoutNotFoundError(DOUYIN.not_found_message())
        assert (
            _parse_failure_reason(exc)
            == "没有找到『抖音博主库』sheet 或『博主ID』表头，请上传灰豚导出的原文件"
        )
        assert _parse_failure_reason(KeyError("x")) == "parse_error:KeyError"
        assert _parse_failure_reason(ValueError("x")) == "parse_error:ValueError"


class TestDuplicateNames:
    def test_no_group_row_numbers_duplicates(self) -> None:
        raw = _book({"抖音博主库": [["博主ID", "点赞", "点赞", "点赞"], ["T1", 1, 2, 3]]})
        rows = _parse_rows(raw, "x.xlsx", DOUYIN)
        assert rows == [(1, {"博主ID": "T1", "点赞": "1", "点赞#2": "2", "点赞#3": "3"})]

    def test_group_row_must_be_directly_above(self) -> None:
        raw = _book({"抖音博主库": [["分组A"], [], ["博主ID", "点赞", "点赞"], ["T1", 1, 2]]})
        rows = _parse_rows(raw, "x.xlsx", DOUYIN)
        assert rows == [(1, {"博主ID": "T1", "点赞": "1", "点赞#2": "2"})]

    def test_same_group_still_duplicate(self) -> None:
        raw = _book({"抖音博主库": [[None, "分组A"], ["博主ID", "点赞", "点赞"], ["T1", 1, 2]]})
        rows = _parse_rows(raw, "x.xlsx", DOUYIN)
        assert rows == [(1, {"博主ID": "T1", "分组A·点赞": "1", "分组A·点赞#2": "2"})]

    def test_rename_helper(self) -> None:
        # 第 0 列向左取不到分组 → 不加前缀；第 2 列向左就近取到「G」，与第 1 列重名 → #2
        assert rename_duplicate_headers(["点赞", "点赞", "点赞", "", ""], ["", "G"]) == [
            "点赞",
            "G·点赞",
            "G·点赞#2",
            "",
            "",
        ]
        assert rename_duplicate_headers(["a", "b"], ["G", "H"]) == ["a", "b"]  # 不重名的不改
        assert rename_duplicate_headers(["x", "x"], None) == ["x", "x#2"]

    def test_without_group_prefix_flag_keeps_legacy_overwrite(self) -> None:
        raw = _book({"抖音博主库": [["分组A"], ["博主ID", "点赞", "点赞"], ["T1", 1, 2]]})
        layout = XlsxLayout(sheet="抖音博主库", header_marker="博主ID")
        assert _parse_rows(raw, "x.xlsx", layout) == [(1, {"博主ID": "T1", "点赞": "2"})]
