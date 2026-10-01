"""U17 集成测试：用户偏好 + 报表导出。

原来这里还测过 bundle_product / bundle_item 的套装拆分，那套在 migration 045 整体
删除了（router 从未注册、生产 0 行），套装改由 goods_main.is_suit 承担，相应测试在
tests/integration/test_goods_crud.py。
"""

from __future__ import annotations

import io
from datetime import date
from typing import Any

import pytest
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.tenancy import tenant_id_ctx
from app.modules.report.bi_service import DEFAULT_BI_LAYOUT
from app.modules.report.export_service import ReportExportService
from app.modules.report.user_preference_service import UserPreferenceService

pytestmark = pytest.mark.asyncio


class TestUserPreference:
    async def test_upsert_and_get_default(
        self,
        session: AsyncSession,
        tenant_a: Any,
        factory: Any,
        admin_role: Any,
    ) -> None:
        tok = tenant_id_ctx.set(tenant_a.id)
        try:
            user = await factory.user(tenant_a, roles=[admin_role])
            svc = UserPreferenceService(session)
            # 无 → 默认
            got = await svc.get_or_default(user.id, "bi_layout", DEFAULT_BI_LAYOUT)
            assert got == DEFAULT_BI_LAYOUT
            # upsert → 回显
            await svc.upsert(user, "bi_layout", {"cards": ["x"]})
            got2 = await svc.get_or_default(user.id, "bi_layout", DEFAULT_BI_LAYOUT)
            assert got2 == {"cards": ["x"]}
        finally:
            tenant_id_ctx.reset(tok)


class TestExport:
    async def test_export_production_xlsx_parseable(
        self,
        session: AsyncSession,
        tenant_a: Any,
    ) -> None:
        tok = tenant_id_ctx.set(tenant_a.id)
        try:
            resp = await ReportExportService(session).export(
                tenant_a.id, "production", (date(2026, 6, 1), date(2026, 6, 30))
            )
            body = b"".join([chunk async for chunk in resp.body_iterator])
            wb = load_workbook(io.BytesIO(body))
            ws = wb.active
            header = [c.value for c in next(ws.iter_rows())]
            assert "商品编码" in header and "含款号" in header and "净投产比" in header
        finally:
            tenant_id_ctx.reset(tok)

    async def test_export_invalid_type(
        self,
        session: AsyncSession,
        tenant_a: Any,
    ) -> None:
        from app.modules.report.exceptions import ReportExportTypeInvalidError

        tok = tenant_id_ctx.set(tenant_a.id)
        try:
            with pytest.raises(ReportExportTypeInvalidError):
                await ReportExportService(session).export(
                    tenant_a.id, "unknown", (date(2026, 6, 1), date(2026, 6, 30))
                )
        finally:
            tenant_id_ctx.reset(tok)
